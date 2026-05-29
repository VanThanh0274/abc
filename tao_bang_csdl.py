from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ---- Page margins ----
section = doc.sections[0]
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3)
section.right_margin  = Cm(2)

# ---- Default font ----
style = doc.styles['Normal']
font  = style.font
font.name = 'Times New Roman'
font.size = Pt(13)

def set_font(run, bold=False, size=13, color=None):
    run.font.name  = 'Times New Roman'
    run.font.size  = Pt(size)
    run.font.bold  = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def heading(text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_font(run, bold=True, size=13)
    return p

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top','left','bottom','right'):
        bd = OxmlElement(f'w:{edge}')
        bd.set(qn('w:val'),   'single')
        bd.set(qn('w:sz'),    '4')
        bd.set(qn('w:space'), '0')
        bd.set(qn('w:color'), '2C3E50')
        tcBorders.append(bd)
    tcPr.append(tcBorders)

HEADER_COLOR = '2C3E50'   # dark navy
ROW_ODD      = 'EBF5FB'   # light blue
ROW_EVEN     = 'FFFFFF'   # white

def add_table(title, headers, rows):
    # Table title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(title)
    set_font(run, bold=True, size=13)

    col_count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # Column widths (cm): Tên trường | Kiểu dữ liệu | Ràng buộc | Mô tả
    widths = [Cm(3.8), Cm(4.2), Cm(3.2), Cm(5.8)]

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        cell.width = widths[i]
        set_cell_bg(cell, HEADER_COLOR)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p2 = cell.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p2.add_run(h)
        set_font(run, bold=True, size=12, color=(255, 255, 255))

    # Data rows
    for idx, row_data in enumerate(rows):
        row_cells = table.rows[idx + 1].cells
        bg = ROW_ODD if idx % 2 == 0 else ROW_EVEN
        for j, val in enumerate(row_data):
            cell = row_cells[j]
            cell.width = widths[j]
            set_cell_bg(cell, bg)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p3 = cell.paragraphs[0]
            p3.alignment = WD_ALIGN_PARAGRAPH.LEFT if j != 2 else WD_ALIGN_PARAGRAPH.CENTER
            is_bold = (j == 0)
            run = p3.add_run(val)
            set_font(run, bold=is_bold, size=12)

    doc.add_paragraph()  # spacing after table

# ========== HEADERS ==========
HEADERS = ['Tên trường', 'Kiểu dữ liệu', 'Ràng buộc', 'Mô tả']

# ============================================================
# TITLE
# ============================================================
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_t = p_title.add_run('MÔ TẢ CẤU TRÚC BẢNG CƠ SỞ DỮ LIỆU')
set_font(run_t, bold=True, size=14)
doc.add_paragraph()

# ============================================================
# 1. NguoiDung
# ============================================================
add_table(
    'Bảng 1: Bảng NguoiDung (Người dùng)',
    HEADERS,
    [
        ('id',        'INT',            'PK, IDENTITY',    'Khóa chính, tự tăng'),
        ('username',  'NVARCHAR(255)',  'UNIQUE, NOT NULL', 'Tên đăng nhập'),
        ('pass',      'NVARCHAR(255)',  'NOT NULL',         'Mật khẩu người dùng'),
        ('ten',       'NVARCHAR(255)',  '',                 'Họ và tên người dùng'),
        ('sdt',       'NVARCHAR(20)',   '',                 'Số điện thoại'),
        ('email',     'NVARCHAR(255)',  '',                 'Địa chỉ email'),
        ('role',      'NVARCHAR(50)',   '',                 'Vai trò: Admin / User'),
        ('trangthai', 'INT',            '',                 'Trạng thái tài khoản (1: Hoạt động, 0: Bị khóa)'),
        ('is_vip',    'INT',            '',                 'Thành viên VIP (1: VIP, 0: Thường)'),
    ]
)

# ============================================================
# 2. Kinhmat
# ============================================================
add_table(
    'Bảng 2: Bảng Kinhmat (Sản phẩm kính mắt)',
    HEADERS,
    [
        ('id',         'INT',           'PK, IDENTITY', 'Khóa chính, tự tăng'),
        ('ten',        'NVARCHAR(255)', 'NOT NULL',     'Tên sản phẩm'),
        ('soluong',    'INT',           '',             'Số lượng tồn kho'),
        ('madanhmuc',  'INT',           'FK',           'Mã danh mục (liên kết bảng DanhMuc)'),
        ('giaban',     'INT',           '',             'Giá bán sản phẩm (VNĐ)'),
        ('gianhap',    'INT',           '',             'Giá nhập sản phẩm (VNĐ)'),
        ('gioithieu',  'NVARCHAR(MAX)', '',             'Đối tượng sử dụng (Nam / Nữ / Unisex)'),
        ('xuatxu',     'NVARCHAR(255)', '',             'Xuất xứ sản phẩm'),
        ('chatlieu',   'NVARCHAR(255)', '',             'Chất liệu gọng kính'),
        ('kieudang',   'NVARCHAR(255)', '',             'Kiểu dáng gọng kính'),
        ('mota',       'NVARCHAR(MAX)', '',             'Mô tả chi tiết sản phẩm'),
        ('anh',        'NVARCHAR(255)', '',             'Đường dẫn ảnh sản phẩm'),
        ('trangthai',  'INT',           '',             'Trạng thái (1: Còn hàng, 0: Hết hàng)'),
    ]
)

# ============================================================
# 3. DanhMuc
# ============================================================
add_table(
    'Bảng 3: Bảng DanhMuc (Danh mục sản phẩm)',
    HEADERS,
    [
        ('ma',         'INT',           'PK, IDENTITY', 'Khóa chính, tự tăng'),
        ('tendanhmuc', 'NVARCHAR(255)', 'NOT NULL',     'Tên danh mục'),
        ('mota',       'NVARCHAR(MAX)', '',             'Mô tả danh mục'),
    ]
)

# ============================================================
# 4. NhaCungCap
# ============================================================
add_table(
    'Bảng 4: Bảng NhaCungCap (Nhà cung cấp)',
    HEADERS,
    [
        ('id',       'INT',           'PK, IDENTITY', 'Khóa chính, tự tăng'),
        ('tenncc',   'NVARCHAR(255)', 'NOT NULL',     'Tên nhà cung cấp'),
        ('sdt',      'NVARCHAR(20)',  '',             'Số điện thoại liên hệ'),
        ('email',    'NVARCHAR(255)', '',             'Địa chỉ email'),
        ('diachi',   'NVARCHAR(500)', '',             'Địa chỉ nhà cung cấp'),
        ('trangthai','INT',           '',             'Trạng thái (1: Đang hợp tác, 0: Ngừng)'),
    ]
)

# ============================================================
# 5. HoaDonBan
# ============================================================
add_table(
    'Bảng 5: Bảng HoaDonBan (Đơn hàng)',
    HEADERS,
    [
        ('mahd',      'INT',           'PK, IDENTITY', 'Mã đơn hàng, tự tăng'),
        ('iduser',    'INT',           'FK',           'Mã người đặt hàng (liên kết bảng NguoiDung)'),
        ('ten',       'NVARCHAR(255)', '',             'Tên người nhận hàng'),
        ('sdt',       'NVARCHAR(20)',  '',             'Số điện thoại người nhận'),
        ('email',     'NVARCHAR(255)', '',             'Email người nhận'),
        ('diachi',    'NVARCHAR(500)', '',             'Địa chỉ giao hàng'),
        ('ghichu',    'NVARCHAR(MAX)', '',             'Ghi chú đơn hàng'),
        ('thoigian',  'DATETIME',      '',             'Thời gian đặt hàng'),
        ('tongtien',  'INT',           '',             'Tổng tiền đơn hàng (VNĐ)'),
        ('tien_giam', 'INT',           '',             'Số tiền được giảm từ khuyến mãi'),
        ('ma_km',     'NVARCHAR(50)',  '',             'Mã khuyến mãi đã áp dụng'),
        ('trangthai', 'NVARCHAR(100)', '',             'Trạng thái đơn hàng'),
    ]
)

# ============================================================
# 6. ChiTietDonHang
# ============================================================
add_table(
    'Bảng 6: Bảng ChiTietDonHang (Chi tiết đơn hàng)',
    HEADERS,
    [
        ('macthd',  'INT', 'PK, IDENTITY', 'Mã chi tiết, tự tăng'),
        ('mahd',    'INT', 'FK',           'Mã đơn hàng (liên kết bảng HoaDonBan)'),
        ('masp',    'INT', 'FK',           'Mã sản phẩm (liên kết bảng Kinhmat)'),
        ('soluong', 'INT', '',             'Số lượng sản phẩm đặt mua'),
        ('giaban',  'INT', '',             'Giá bán tại thời điểm đặt hàng'),
    ]
)

# ============================================================
# 7. Khuyenmai
# ============================================================
add_table(
    'Bảng 7: Bảng Khuyenmai (Khuyến mãi)',
    HEADERS,
    [
        ('id',                 'INT',           'PK, IDENTITY',    'Khóa chính, tự tăng'),
        ('ma_km',              'NVARCHAR(50)',  'UNIQUE, NOT NULL', 'Mã code khuyến mãi'),
        ('ten_km',             'NVARCHAR(255)', '',                'Tên chương trình khuyến mãi'),
        ('mota',               'NVARCHAR(MAX)', '',                'Mô tả khuyến mãi'),
        ('loai_km',            'INT',           '',                'Loại giảm (1: Giảm %, 0: Giảm tiền trực tiếp)'),
        ('giatri_km',          'INT',           '',                'Giá trị giảm giá'),
        ('dieukien_toithieu',  'INT',           '',                'Giá trị đơn hàng tối thiểu để áp dụng'),
        ('ngaybatdau',         'DATETIME',      '',                'Ngày bắt đầu hiệu lực'),
        ('ngayketthuc',        'DATETIME',      '',                'Ngày kết thúc hiệu lực'),
        ('is_vip_only',        'INT',           '',                'Chỉ dành cho VIP (1: VIP only, 0: Tất cả)'),
        ('trangthai',          'INT',           '',                'Trạng thái (1: Đang áp dụng, 0: Hết hạn)'),
    ]
)

# ============================================================
# 8. HoaDonNhap
# ============================================================
add_table(
    'Bảng 8: Bảng HoaDonNhap (Hóa đơn nhập kho)',
    HEADERS,
    [
        ('mahdn',      'INT',           'PK, IDENTITY', 'Mã hóa đơn nhập, tự tăng'),
        ('mancc',      'INT',           'FK',           'Mã nhà cung cấp (liên kết bảng NhaCungCap)'),
        ('ngaynhap',   'DATETIME',      '',             'Ngày thực hiện nhập kho'),
        ('tongtien',   'INT',           '',             'Tổng tiền hóa đơn nhập (VNĐ)'),
        ('nguoinhap',  'NVARCHAR(255)', '',             'Người thực hiện nhập kho'),
        ('ghichu',     'NVARCHAR(MAX)', '',             'Ghi chú hóa đơn'),
        ('trangthai',  'INT',           '',             'Trạng thái (0: Nháp, 1: Đã nhập kho)'),
    ]
)

# ============================================================
# 9. ChiTietHoaDonNhap
# ============================================================
add_table(
    'Bảng 9: Bảng ChiTietHoaDonNhap (Chi tiết nhập kho)',
    HEADERS,
    [
        ('macthdn', 'INT', 'PK, IDENTITY', 'Mã chi tiết nhập, tự tăng'),
        ('mahdn',   'INT', 'FK',           'Mã hóa đơn nhập (liên kết bảng HoaDonNhap)'),
        ('masp',    'INT', 'FK',           'Mã sản phẩm (liên kết bảng Kinhmat)'),
        ('soluong', 'INT', '',             'Số lượng nhập kho'),
        ('gianhap', 'INT', '',             'Giá nhập tại thời điểm nhập kho'),
    ]
)

# ============================================================
# 10. LichSuGia
# ============================================================
add_table(
    'Bảng 10: Bảng LichSuGia (Lịch sử giá)',
    HEADERS,
    [
        ('id',          'INT',      'PK, IDENTITY', 'Khóa chính, tự tăng'),
        ('masp',        'INT',      'FK',           'Mã sản phẩm (liên kết bảng Kinhmat)'),
        ('gianhap',     'INT',      '',             'Giá nhập tại thời điểm áp dụng'),
        ('giaban',      'INT',      '',             'Giá bán tại thời điểm áp dụng'),
        ('ngayapdung',  'DATETIME', '',             'Ngày bắt đầu áp dụng mức giá'),
        ('ngayketthuc', 'DATETIME', '',             'Ngày kết thúc áp dụng mức giá (NULL = đang hiệu lực)'),
        ('trangthai',   'INT',      '',             'Trạng thái (1: Đang áp dụng, 0: Đã hết hạn)'),
    ]
)

# ============================================================
# 11. Blog
# ============================================================
add_table(
    'Bảng 11: Bảng Blog (Bài viết)',
    HEADERS,
    [
        ('id',          'INT',           'PK, IDENTITY', 'Khóa chính, tự tăng'),
        ('tieude',      'NVARCHAR(500)', '',             'Tiêu đề bài viết'),
        ('noidung',     'NVARCHAR(MAX)', '',             'Nội dung bài viết (định dạng Markdown)'),
        ('tomtat',      'NVARCHAR(MAX)', '',             'Tóm tắt bài viết'),
        ('anh',         'NVARCHAR(255)', '',             'Đường dẫn ảnh đại diện'),
        ('madanhmuc',   'INT',           'FK',           'Mã danh mục blog'),
        ('tacgia',      'NVARCHAR(255)', '',             'Tên tác giả bài viết'),
        ('luotxem',     'INT',           '',             'Số lượt xem bài viết'),
        ('trangthai',   'INT',           '',             'Trạng thái (1: Hiển thị, 0: Ẩn)'),
        ('ngaytao',     'DATETIME',      '',             'Ngày tạo bài viết'),
        ('ngaycapnhat', 'DATETIME',      '',             'Ngày cập nhật gần nhất'),
    ]
)

# ============================================================
# 12. LichSuChatbot
# ============================================================
add_table(
    'Bảng 12: Bảng LichSuChatbot (Lịch sử trò chuyện AI)',
    HEADERS,
    [
        ('id',       'INT',           'PK, IDENTITY', 'Khóa chính, tự tăng'),
        ('iduser',   'INT',           'FK',           'Mã người dùng (liên kết bảng NguoiDung)'),
        ('role',     'NVARCHAR(20)',  '',             'Vai trò tin nhắn (user / model)'),
        ('content',  'NVARCHAR(MAX)', '',             'Nội dung tin nhắn'),
        ('thoigian', 'DATETIME',      '',             'Thời gian gửi tin nhắn'),
    ]
)

# ============================================================
# SAVE
# ============================================================
output_path = r'd:\DATN_KinhMat\abc\Mo_ta_CSDL_KinhMat.docx'
doc.save(output_path)
print(f'[OK] File da duoc tao: {output_path}')
